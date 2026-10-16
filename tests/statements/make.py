# Генерирует образцы выписок разных банков (выдуманные данные) + ожидаемые итоги.
import json, random, csv, io, os
from datetime import date, timedelta
import openpyxl
random.seed(7)
HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "out"); os.makedirs(OUT, exist_ok=True)
MER = [("PYATEROCHKA 1234","Супермаркеты"),("MAGNIT MM KRASNYJ","Супермаркеты"),("АЗС Татнефть №45","Топливо"),("LUKOIL AZS 116","Топливо"),
       ("APTEKA VITA","Аптеки"),("KFC CHELNY","Фастфуд"),("LEROY MERLIN","Дом и ремонт"),("OZON.RU","Маркетплейсы"),("MTS","Связь"),
       ("Яндекс Плюс","Подписки"),("AVTOZAPCHASTI 24","Автоуслуги"),("WILDBERRIES","Маркетплейсы")]
ops = []
d0 = date(2026, 9, 1)
for i in range(46):
    d = d0 + timedelta(days=i * 29 // 46)
    if i in (5, 30):
        ops.append(dict(date=d, type="inc", amount=80000.00 if i == 5 else 45500.50, desc="Зарплата ООО Кентавр" if i == 5 else "Перевод от Иван И.", cat="Зарплата" if i == 5 else "Переводы"))
    elif i == 18:
        ops.append(dict(date=d, type="inc", amount=356.70, desc="Возврат OZON.RU", cat="Маркетплейсы"))
    else:
        m, c = MER[i % len(MER)]
        amt = round(random.choice([random.uniform(90, 990), random.uniform(1000, 9000), random.uniform(10000, 23000)]), 2)
        ops.append(dict(date=d, type="exp", amount=amt, desc=m, cat=c))
# одна и та же покупка дважды в день — обе должны остаться
ops.append(dict(ops[1])); ops.sort(key=lambda o: o["date"])
exp = round(sum(o["amount"] for o in ops if o["type"] == "exp"), 2)
inc = round(sum(o["amount"] for o in ops if o["type"] == "inc"), 2)
def rus(v, sep=" ", dec=","):
    s = f"{v:,.2f}".replace(",", "X").replace(".", dec).replace("X", sep); return s
dd = lambda d: d.strftime("%d.%m.%Y")
expected = dict(count=len(ops), exp=exp, inc=inc, first=str(ops[0]["date"]), last=str(ops[-1]["date"]))
CSS = "<style>body{font-family:'DejaVu Sans',Arial;font-size:9pt}table{border-collapse:collapse;width:100%}td,th{border-bottom:1px solid #ccc;padding:3px 4px;text-align:left;vertical-align:top}th{font-size:8pt}</style>"
def page(title, intro, head, rows):
    return f"<!doctype html><meta charset=utf-8>{CSS}<h3>{title}</h3>{intro}<table><thead><tr>{''.join('<th>'+h+'</th>' for h in head)}</tr></thead><tbody>{''.join(rows)}</tbody></table><p>Итого списаний: {rus(exp)}</p><p>Итого поступлений: {rus(inc)}</p>"
# --- Сбер: строка 1 — дата, время, код, категория, сумма (доход с +), остаток; строка 2 — дата и описание
bal = 12000.00; rows = []
for o in ops:
    bal = bal + (o["amount"] if o["type"] == "inc" else -o["amount"])
    sign = "+" if o["type"] == "inc" else ""
    rows.append(f"<tr><td>{dd(o['date'])} 13:45<br>{dd(o['date'])}</td><td>{random.randint(100000,999999)}</td><td>{o['cat']}<br>{o['desc']}. Операция по карте ****1234</td><td>{sign}{rus(o['amount'])}</td><td>{rus(bal)}</td></tr>")
open(f"{OUT}/sber.html","w").write(page("Выписка по счёту дебетовой карты", f"<p>Остаток на {dd(d0)}: 12 000,00</p>", ["ДАТА ОПЕРАЦИИ (МСК)","КОД АВТОРИЗАЦИИ","КАТЕГОРИЯ","СУММА В ВАЛЮТЕ СЧЁТА","ОСТАТОК СРЕДСТВ"], rows))
# --- Т-Банк PDF
rows = [f"<tr><td>{dd(o['date'])}<br>12:0{i%10}</td><td>{dd(o['date'])}</td><td>{'+' if o['type']=='inc' else '-'}{rus(o['amount'])} ₽</td><td>{'+' if o['type']=='inc' else '-'}{rus(o['amount'])} ₽</td><td>{'Оплата в ' if o['type']=='exp' else ''}{o['desc']}</td><td>1234</td></tr>" for i,o in enumerate(ops)]
open(f"{OUT}/tbank.html","w").write(page("Справка о движении средств", "<p>АО «ТБанк»</p>", ["Дата и время операции","Дата списания","Сумма в валюте операции","Сумма операции в валюте карты","Описание операции","Номер карты"], rows))
# --- Альфа PDF: расход со знаком минус, доход без знака
rows = [f"<tr><td>{dd(o['date'])}</td><td>CRD_{i:04d}</td><td>{o['desc']}</td><td>{'-' if o['type']=='exp' else ''}{rus(o['amount'])} RUR</td></tr>" for i,o in enumerate(ops)]
open(f"{OUT}/alfa.html","w").write(page("Выписка по счёту", "<p>Входящий остаток 12 000,00 RUR</p>", ["Дата проводки","Код","Описание","Сумма в валюте счета"], rows))
# --- ВТБ PDF: точка как разделитель копеек
rows = [f"<tr><td>{dd(o['date'])} 10:11:12</td><td>{dd(o['date'])}</td><td>{'-' if o['type']=='exp' else '+'}{rus(o['amount'], sep=' ', dec='.')} RUB</td><td>{'-' if o['type']=='exp' else '+'}{rus(o['amount'], sep=' ', dec='.')} RUB</td><td>{o['desc']}</td></tr>" for o in ops]
open(f"{OUT}/vtb.html","w").write(page("Выписка по мастер-счёту", "", ["Дата операции","Дата обработки","Сумма операции","Сумма в валюте счёта","Основание операции"], rows))
# --- Т-Банк CSV (windows-1251, ;) + одна отклонённая операция, которой не должно быть
buf = io.StringIO(); w = csv.writer(buf, delimiter=";", quoting=csv.QUOTE_ALL, lineterminator="\r\n")
w.writerow(["Дата операции","Дата платежа","Номер карты","Статус","Сумма операции","Валюта операции","Сумма платежа","Валюта платежа","Кэшбэк","Категория","MCC","Описание","Бонусы (включая кэшбэк)"])
for o in ops:
    v = ("-" if o["type"] == "exp" else "") + f"{o['amount']:.2f}".replace(".", ",")
    w.writerow([o["date"].strftime("%d.%m.%Y 13:45:11"), dd(o["date"]), "*1234", "OK", v, "RUB", v, "RUB", "", o["cat"], "5411", o["desc"], "0,00"])
w.writerow(["15.09.2026 10:00:00","15.09.2026","*1234","FAILED","-99999,00","RUB","-99999,00","RUB","","Супермаркеты","5411","ОТКЛОНЕНО","0,00"])
open(f"{OUT}/tbank.csv","wb").write(buf.getvalue().encode("cp1251"))
# --- Альфа CSV (utf-8, запятая, сумма без знака + колонка «Тип»)
buf = io.StringIO(); w = csv.writer(buf, delimiter=",", lineterminator="\n")
w.writerow(["Дата операции","Дата проводки","Код авторизации","Описание","Категория","Сумма в валюте счета","Статус","Тип"])
for o in ops:
    w.writerow([dd(o["date"]), dd(o["date"]), "A1B2C3", o["desc"], o["cat"], f"{o['amount']:.2f}", "Выполнен", "Списание" if o["type"] == "exp" else "Пополнение"])
open(f"{OUT}/alfa.csv","w",encoding="utf-8").write(buf.getvalue())
# --- ВТБ Excel: шапка, даты как даты Excel, суммы числами со знаком
wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Выписка"
ws.append(["Выписка по счёту 40817810..."]); ws.append(["Период: 01.09.2026 – 30.09.2026"]); ws.append([])
ws.append(["Дата операции","Дата обработки","Сумма операции","Валюта","Сумма в валюте счёта","Описание операции"])
for o in ops:
    v = o["amount"] if o["type"] == "inc" else -o["amount"]
    ws.append([o["date"], o["date"], v, "RUB", v, o["desc"]])
    ws.cell(ws.max_row, 1).number_format = "DD.MM.YYYY"; ws.cell(ws.max_row, 2).number_format = "DD.MM.YYYY"
ws.append([]); ws.append(["Итого", None, None, None, round(inc - exp, 2)])
wb.save(f"{OUT}/vtb.xlsx")
# --- Excel «Приход / Расход» (как у многих банков и 1С)
wb = openpyxl.Workbook(); ws = wb.active
ws.append(["Дата","Назначение платежа","Приход","Расход"])
for o in ops:
    ws.append([dd(o["date"]), o["desc"], o["amount"] if o["type"] == "inc" else None, o["amount"] if o["type"] == "exp" else None])
wb.save(f"{OUT}/prihod-rashod.xlsx")
json.dump(expected, open(f"{OUT}/expected.json", "w"), ensure_ascii=False, indent=1)
print(expected)
